from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Iterable

from .models import FULL_VARIABLES
from .rubric import load_questions
from .scoring import normalize_answer, score_case

INSTRUMENT_SHEET = "5-Variable Instrument"
CONVERGENCE_SHEET = "Convergence Map"
EVENT_ROWS = range(12, 19)
VARIABLE_ANSWER_CELLS: dict[str, tuple[str, ...]] = {
    "G": ("B26", "B27", "B28"),
    "T": ("B37", "B38", "B39", "B40", "B41"),
    "K": ("B50", "B51", "B52", "B53"),
    "F": ("B62", "B63", "B64", "B65"),
    "R": ("B73", "B74", "B75", "B76"),
}
VARIABLE_SCORE_CELLS = {"G": "C29", "T": "C42", "K": "C54", "F": "C66", "R": "C77"}
SUMMARY_CELLS = {"direction": "C82", "veto": "C83", "weakest_factor": "C84", "unknown_count": "C85"}
WORKBOOK_PATTERNS = ("*.xlsx", "*.xlsm", "*.xltx", "*.xltm")

_EVENT_ALIASES = {
    "case_id": ("case id", "case_id", "id"),
    "actor": ("actor", "agent"),
    "action": ("action", "act"),
    "target": ("target", "object"),
    "affected_parties": ("affected", "parties", "affected parties"),
    "claimed_objective": ("objective", "claimed objective", "intent"),
    "mechanism": ("mechanism", "means"),
    "cost_bearers": ("cost bearers", "cost_bearers", "cost"),
    "beneficiaries": ("beneficiaries", "beneficiary"),
    "start_date": ("start", "start date"),
    "end_date": ("end", "end date"),
    "measured_outcomes": ("outcome", "measured outcome", "measured outcomes"),
    "counterfactual": ("counterfactual", "baseline"),
}
_DEFAULT_EVENT_ORDER = ("case_id", "actor", "action", "target", "affected_parties", "claimed_objective", "mechanism")
_LIST_FIELDS = {"affected_parties", "cost_bearers", "beneficiaries", "measured_outcomes"}


def find_workbook(root: Path = Path.cwd()) -> Path:
    candidates = [p for pattern in WORKBOOK_PATTERNS for p in root.glob(pattern) if not p.name.startswith("~$")]
    if not candidates:
        raise FileNotFoundError(f"No Excel workbook found in {root}")
    if len(candidates) == 1:
        return candidates[0]
    audited = [p for p in candidates if "AUDITED" in p.stem.upper()]
    return sorted(audited or candidates, key=lambda p: (p.stat().st_mtime, p.name), reverse=True)[0]


def _require_openpyxl() -> Any:
    import openpyxl

    return openpyxl


def _string(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _split_list(value: Any) -> list[str]:
    text = _string(value)
    if not text:
        return []
    return [item.strip() for chunk in text.split(";") for item in chunk.split(",") if item.strip()]


def _event_key(label: str, fallback_index: int) -> str:
    token = label.strip().lower().replace(":", "")
    for key, aliases in _EVENT_ALIASES.items():
        if any(alias in token for alias in aliases):
            return key
    return _DEFAULT_EVENT_ORDER[fallback_index]


def _read_event_fields(ws: Any) -> dict[str, Any]:
    fields: dict[str, Any] = {}
    for index, row in enumerate(EVENT_ROWS):
        key = _event_key(_string(ws[f"A{row}"].value), index)
        value = ws[f"B{row}"].value
        fields[key] = _split_list(value) if key in _LIST_FIELDS else _string(value)
    return fields


def _question_ids_by_variable() -> dict[str, list[str]]:
    grouped: dict[str, list[str]] = {name: [] for name in FULL_VARIABLES}
    for qid, question in load_questions().items():
        grouped.setdefault(question.variable, []).append(qid)
    return grouped


def _answer_from_cell(value: Any) -> str:
    text = _string(value)
    if not text:
        return "UNKNOWN"
    return normalize_answer(text).value


def case_from_workbook(path: Path | None = None) -> dict[str, Any]:
    openpyxl = _require_openpyxl()
    workbook_path = path or find_workbook()
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        ws = wb[INSTRUMENT_SHEET]
        event = _read_event_fields(ws)
        qids = _question_ids_by_variable()
        variables: dict[str, dict[str, list[dict[str, Any]]]] = {}
        for variable, cells in VARIABLE_ANSWER_CELLS.items():
            ids = qids.get(variable, [])
            variables[variable] = {"answers": [
                {"question_id": ids[i] if i < len(ids) else f"{variable}{i + 1}", "answer": _answer_from_cell(ws[cell].value)}
                for i, cell in enumerate(cells)
            ]}
    finally:
        wb.close()
    case = {
        "case_id": event.get("case_id") or "excel-workbook-case",
        "actor": event.get("actor") or "UNKNOWN",
        "action": event.get("action") or "UNKNOWN",
        "target": event.get("target") or "UNKNOWN",
        "affected_parties": event.get("affected_parties") or ["UNKNOWN"],
        "claimed_objective": event.get("claimed_objective") or "UNKNOWN",
        "mechanism": event.get("mechanism") or "UNKNOWN",
        "cost_bearers": event.get("cost_bearers") or event.get("affected_parties") or ["UNKNOWN"],
        "beneficiaries": event.get("beneficiaries") or [event.get("actor") or "UNKNOWN"],
        "start_date": event.get("start_date") or "UNKNOWN",
        "end_date": event.get("end_date") or "UNKNOWN",
        "measured_outcomes": event.get("measured_outcomes") or ["UNKNOWN"],
        "counterfactual": event.get("counterfactual") or "UNKNOWN",
        "variables": variables,
    }
    return case


def _unknown_count(result: dict[str, Any]) -> int:
    return sum(1 for data in result.get("variables", {}).values() if data.get("raw_score") is None)


def write_scores(path: Path, result: dict[str, Any]) -> None:
    openpyxl = _require_openpyxl()
    wb = openpyxl.load_workbook(path)
    try:
        ws = wb[INSTRUMENT_SHEET]
        for variable, cell in VARIABLE_SCORE_CELLS.items():
            ws[cell] = result.get("variables", {}).get(variable, {}).get("rounded_score")
        direction = result.get("direction", {})
        ws[SUMMARY_CELLS["direction"]] = direction.get("label")
        ws[SUMMARY_CELLS["veto"]] = result.get("veto", {}).get("variable")
        ws[SUMMARY_CELLS["weakest_factor"]] = result.get("veto", {}).get("variable")
        ws[SUMMARY_CELLS["unknown_count"]] = _unknown_count(result)
        wb.save(path)
    finally:
        wb.close()


def read_convergence(path: Path | None = None) -> list[dict[str, Any]]:
    openpyxl = _require_openpyxl()
    workbook_path = path or find_workbook()
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        ws = wb[CONVERGENCE_SHEET]
        headers = [_string(cell.value) for cell in ws[1] if _string(cell.value)]
        rows: list[dict[str, Any]] = []
        for values in ws.iter_rows(min_row=2, max_col=len(headers), values_only=True):
            if not any(value is not None and _string(value) for value in values):
                continue
            rows.append({headers[i]: values[i] for i in range(len(headers))})
        return rows
    finally:
        wb.close()


def write_convergence(path: Path, rows: Iterable[dict[str, Any]], start_row: int = 2) -> None:
    openpyxl = _require_openpyxl()
    wb = openpyxl.load_workbook(path)
    try:
        ws = wb[CONVERGENCE_SHEET]
        headers = [_string(cell.value) for cell in ws[1] if _string(cell.value)]
        for row_offset, row in enumerate(rows, start=start_row):
            for col_index, header in enumerate(headers, start=1):
                if header in row:
                    ws.cell(row=row_offset, column=col_index).value = row[header]
        wb.save(path)
    finally:
        wb.close()


def sync_workbook(path: Path | None = None) -> dict[str, Any]:
    workbook_path = path or find_workbook()
    case = case_from_workbook(workbook_path)
    result = score_case(case)
    if result.get("status") == "SCORED":
        write_scores(workbook_path, result)
    return {"workbook": str(workbook_path), "case": case, "result": result}


def main() -> None:
    parser = argparse.ArgumentParser(description="Synchronize the Meta-Argument Excel workbook with the Python scoring engine")
    parser.add_argument("workbook", nargs="?", type=Path, help="Workbook path; defaults to the Excel file in the repo root")
    parser.add_argument("--case-json", type=Path, help="Optional path to write the constructed case JSON")
    parser.add_argument("--result-json", type=Path, help="Optional path to write the scoring result JSON")
    args = parser.parse_args()
    payload = sync_workbook(args.workbook)
    if args.case_json:
        args.case_json.write_text(json.dumps(payload["case"], indent=2, ensure_ascii=False), encoding="utf-8")
    if args.result_json:
        args.result_json.write_text(json.dumps(payload["result"], indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(payload["result"], indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
